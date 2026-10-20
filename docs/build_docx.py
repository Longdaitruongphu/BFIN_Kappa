#!/usr/bin/env python3
"""Build a polished DOCX from the controlled Markdown used in this docs folder.

Usage:
    python build_docx.py input.md output.docx "Header text"
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2E, 0x54, 0x96)
DARK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BLUE = "EAF1FB"
ROW_ALT = "F3F6FB"
HEADER_FILL = "1F3864"
BORDER = "B4C6E7"

BODY_FONT = "Arial"
CODE_FONT = "Courier New"


def set_run_font(run, name=BODY_FONT, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "vi-VN")
    rPr.append(lang)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for tag, value in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def set_table_borders(table, color=BORDER, size=4):
    if isinstance(color, RGBColor):
        color = str(color)
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)
        borders.append(element)
    tblPr.append(borders)


def mark_header_row(row):
    trPr = row._tr.get_or_add_trPr()
    cantSplit = OxmlElement("w:cantSplit")
    trPr.append(cantSplit)
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    cantSplit = OxmlElement("w:cantSplit")
    trPr.append(cantSplit)


def keep_table_together(table):
    """Ask Word/LibreOffice not to break a small table across pages."""
    for row in table.rows[:-1]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True


def add_bottom_border(paragraph, color=BORDER, size=6):
    if isinstance(color, RGBColor):
        color = str(color)
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_page_number_footer(section, header_text, document):
    # Header
    header = section.header
    hp = header.paragraphs[0]
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = hp.add_run(header_text)
    set_run_font(run, size=8, color=GREY)
    add_bottom_border(hp, color=BORDER, size=4)

    # Footer
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    r1 = fp.add_run("Trang ")
    set_run_font(r1, size=8, color=GREY)

    # A PAGE field with a cached result so viewers that do not update fields
    # still show a page number. Word updates it on open/print.
    run_begin = fp.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    run_begin._r.append(fld_begin)

    run_instr = fp.add_run()
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = " PAGE "
    run_instr._r.append(instr_el)

    run_sep = fp.add_run()
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    run_sep._r.append(fld_sep)

    run_cached = fp.add_run("1")
    set_run_font(run_cached, size=8, color=GREY)

    run_end = fp.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run_end._r.append(fld_end)

    # Ask Word to refresh fields on open.
    settings = document.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*]+?\*)")


def add_inline(paragraph, text, size=10.3, color=DARK, bold=False):
    parts = INLINE_RE.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, name=CODE_FONT, size=size - 0.8, color=RGBColor(0x1F, 0x38, 0x64))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=size, italic=True, color=color)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, bold=bold, color=color)


def add_heading(doc, text, level, page_break=False):
    if page_break:
        doc.add_page_break()
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, size=19, bold=True, color=NAVY)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(10)
    elif level == 2:
        set_run_font(run, size=14.5, bold=True, color=BLUE)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(5)
    elif level == 3:
        set_run_font(run, size=11.5, bold=True, color=NAVY)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
    else:
        set_run_font(run, size=10.3, bold=True, color=DARK)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
    return p


def add_bullet(doc, text, ordered=False):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.65)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    p.paragraph_format.space_after = Pt(2)
    prefix = "•  " if not ordered else ""
    if prefix:
        run = p.add_run(prefix)
        set_run_font(run, size=10.3, color=BLUE, bold=True)
    add_inline(p, text, size=10.3)
    return p


def add_callout(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color=BLUE, size=6)
    cell = table.cell(0, 0)
    cell.width = Cm(17.2)
    shade_cell(cell, LIGHT_BLUE)
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_inline(p, text, size=9.8, color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, rows):
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    normalized = []
    for r in rows:
        r = list(r) + [""] * (ncols - len(r))
        normalized.append(r)

    # Estimate column weights from content length and header length.
    max_lens = []
    for j in range(ncols):
        content_len = max(len(re.sub(r"[*`]", "", normalized[i][j])) for i in range(len(normalized)))
        header_len = len(re.sub(r"[*`]", "", normalized[0][j]))
        max_len = max(content_len, int(header_len * 0.9), 4)
        max_lens.append(min(42, max_len))
    total = sum(max_lens)
    available_cm = 17.2
    widths = [max(1.4, available_cm * w / total) for w in max_lens]
    scale = available_cm / sum(widths)
    widths = [w * scale for w in widths]

    table = doc.add_table(rows=len(normalized), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False
    set_table_borders(table)

    # Fixed layout and explicit grid widths: without these, Word/LibreOffice
    # can fall back to equal-width columns and squeeze text columns.
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblLayout = OxmlElement("w:tblLayout")
    tblLayout.set(qn("w:type"), "fixed")
    tblPr.append(tblLayout)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), str(int(sum(widths) * 567)))
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)

    for j, width in enumerate(widths):
        table.columns[j].width = Cm(width)
        grid_col = tbl.tblGrid.gridCol_lst[j]
        grid_col.set(qn("w:w"), str(int(width * 567)))

    font_size = 8.3 if ncols <= 4 else 7.7
    header_size = font_size

    for i, row in enumerate(normalized):
        for j, text in enumerate(row):
            cell = table.cell(i, j)
            cell.width = Cm(widths[j])
            set_cell_margins(cell, top=45, bottom=45, left=65, right=65)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if i == 0:
                shade_cell(cell, HEADER_FILL)
                add_inline(p, text, size=header_size, color=WHITE, bold=True)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                if i % 2 == 0:
                    shade_cell(cell, ROW_ALT)
                add_inline(p, text, size=font_size, color=DARK)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if i == 0:
            mark_header_row(table.rows[i])
        else:
            prevent_row_split(table.rows[i])

    if len(normalized) <= 12:
        keep_table_together(table)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def is_separator_row(line):
    s = line.strip()
    if not s.startswith("|"):
        return False
    cells = [c.strip() for c in s.strip("|").split("|")]
    return all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c != "")


def parse_table_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def build(input_path: Path, output_path: Path, header_text: str, simple: bool = False):
    lines = input_path.read_text(encoding="utf-8").splitlines()

    doc = Document()

    # Page setup
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    # Base styles
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.3)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.08

    for style_name in ("Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        st = doc.styles[style_name]
        st.font.name = BODY_FONT

    add_page_number_footer(section, header_text, doc)

    # Title page
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("BFIN WS2026  ·  GROUP ASSIGNMENT")
    set_run_font(r, size=10, bold=True, color=BLUE)
    p.paragraph_format.space_after = Pt(10)

    p = doc.add_paragraph()
    r = p.add_run("KAPPA FOUNDATION")
    set_run_font(r, size=30, bold=True, color=NAVY)
    p.paragraph_format.space_after = Pt(2)

    p = doc.add_paragraph()
    r = p.add_run("PHÂN CÔNG CÔNG VIỆC NHÓM 4 NGƯỜI")
    set_run_font(r, size=17, bold=True, color=BLUE)
    p.paragraph_format.space_after = Pt(6)

    p = doc.add_paragraph()
    r = p.add_run("Từ đề bài PDF đến sản phẩm nộp 29/10/2026")
    set_run_font(r, size=12, italic=True, color=GREY)
    p.paragraph_format.space_after = Pt(14)

    p = doc.add_paragraph()
    add_bottom_border(p, color=BLUE, size=8)
    p.paragraph_format.space_after = Pt(16)

    meta = [
        ("Prepared for", "Minh — Python Lead & Overall Lead"),
        ("Version", "1.0 — 18/10/2026"),
        ("Status", "Kế hoạch làm việc nội bộ, không phải bằng chứng công việc đã hoàn thành"),
        ("Scope", "Đề gốc 31/12/2022; revised mandate 23/10/2026 chưa có"),
    ]
    table = doc.add_table(rows=len(meta), cols=2)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, (k, v) in enumerate(meta):
        table.cell(i, 0).width = Cm(3.8)
        table.cell(i, 1).width = Cm(12.8)
        p0 = table.cell(i, 0).paragraphs[0]
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(k)
        set_run_font(r0, size=9, bold=True, color=BLUE)
        p1 = table.cell(i, 1).paragraphs[0]
        p1.paragraph_format.space_after = Pt(2)
        add_inline(p1, v, size=9.5, color=DARK)

    doc.add_paragraph()
    add_callout(
        doc,
        "Tài liệu này là **kế hoạch điều hành**, không phải contribution statement. "
        "Chỉ điền tên, việc đã làm, người kiểm và critique sau khi công việc thực sự hoàn thành.",
    )

    # Skip original title metadata until the first horizontal rule after Status.
    i = 0
    while i < len(lines):
        if lines[i].strip() == "---":
            i += 1
            break
        i += 1

    in_table = False
    table_rows = []
    paragraph_buffer = []
    first_part = True

    def flush_paragraph():
        nonlocal paragraph_buffer
        if not paragraph_buffer:
            return
        text = " ".join(x.strip() for x in paragraph_buffer).strip()
        paragraph_buffer = []
        if not text:
            return
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(p, text, size=10.3)

    def flush_table():
        nonlocal table_rows
        if table_rows:
            add_table(doc, table_rows)
            table_rows = []

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if not stripped:
            flush_paragraph()
            i += 1
            continue

        # Table block
        if stripped.startswith("|") and stripped.endswith("|"):
            flush_paragraph()
            table_rows = []
            while i < len(lines):
                candidate = lines[i].strip()
                if not (candidate.startswith("|") and candidate.endswith("|")):
                    break
                if is_separator_row(candidate):
                    i += 1
                    continue
                table_rows.append(parse_table_row(candidate))
                i += 1
            flush_table()
            continue

        # Horizontal rule
        if stripped == "---":
            flush_paragraph()
            i += 1
            continue

        # Headings
        if stripped.startswith("#"):
            flush_paragraph()
            m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
            if m:
                level = len(m.group(1))
                text = m.group(2).strip()
                if level == 1:
                    do_break = not text.startswith("PHẦN II")
                    add_heading(doc, text, 1, page_break=do_break)
                    first_part = False
                elif text.strip().upper() == "MỤC LỤC":
                    add_heading(doc, text, level, page_break=True)
                else:
                    add_heading(doc, text, level)
                i += 1
                continue

        # Blockquote
        if stripped.startswith(">"):
            flush_paragraph()
            add_callout(doc, stripped.lstrip(">").strip())
            i += 1
            continue

        # Bullets
        if stripped.startswith("- "):
            flush_paragraph()
            add_bullet(doc, stripped[2:].strip(), ordered=False)
            i += 1
            continue

        # Numbered list
        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            flush_paragraph()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.85)
            p.paragraph_format.first_line_indent = Cm(-0.55)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(f"{m.group(1)}.  ")
            set_run_font(r, size=10.3, bold=True, color=BLUE)
            add_inline(p, m.group(2), size=10.3)
            i += 1
            continue

        # Plain paragraph text (accumulate wrapped lines)
        paragraph_buffer.append(stripped)
        i += 1

    flush_paragraph()
    flush_table()

    # Core properties
    doc.core_properties.title = "Kappa Foundation — Phân công công việc nhóm 4 người"
    doc.core_properties.subject = "BFIN WS2026 Group Assignment"
    doc.core_properties.author = "Minh"
    doc.core_properties.comments = "Internal working plan; not evidence of completed work."

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(f"Wrote {output_path}")


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        raise SystemExit(1)
    input_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])
    header_text = sys.argv[3] if len(sys.argv) > 3 else "Kappa Foundation · Phân công nhóm 4"
    build(input_path, output_path, header_text)


if __name__ == "__main__":
    main()

