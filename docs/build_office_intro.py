from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2E, 0x54, 0x96)
GREY = RGBColor(0x59, 0x59, 0x59)
DARK = RGBColor(0x22, 0x22, 0x22)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEADER_FILL = "1F3864"
ALT_FILL = "F3F6FB"
BORDER = "B4C6E7"
FONT = "Arial"


def set_font(run, size=9.0, bold=False, italic=False, color=DARK):
    run.font.name = FONT
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def cell_margins(cell, top=30, bottom=30, left=70, right=70):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for tag, value in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def borders(table):
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), BORDER)
        tblBorders.append(el)
    tblPr.append(tblBorders)


def header_row(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)


def add_text(p, text, size=9.0, color=DARK, bold=False):
    parts = text.split("**")
    for i, part in enumerate(parts):
        if not part:
            continue
        r = p.add_run(part)
        set_font(r, size=size, color=color, bold=bold or (i % 2 == 1))


def heading(doc, text, size=10.5, before=4, after=1.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(r, size=size, bold=True, color=BLUE)
    return p


def bullet(doc, text, size=9.0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.42)
    p.paragraph_format.first_line_indent = Cm(-0.22)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("•  ")
    set_font(r, size=size, color=BLUE, bold=True)
    add_text(p, text, size=size)


def table(doc, rows, widths, font_size=8.4):
    table = doc.add_table(rows=len(rows), cols=len(widths))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False
    borders(table)
    tbl = table._tbl
    tblPr = tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), str(int(sum(widths) * 567)))
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)
    for j, w in enumerate(widths):
        table.columns[j].width = Cm(w)
        tbl.tblGrid.gridCol_lst[j].set(qn("w:w"), str(int(w * 567)))
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            cell = table.cell(i, j)
            cell.width = Cm(widths[j])
            cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if i == 0:
                shade(cell, HEADER_FILL)
                add_text(p, text, size=font_size, color=WHITE, bold=True)
            else:
                if i % 2 == 0:
                    shade(cell, ALT_FILL)
                add_text(p, text, size=font_size, color=DARK)
        if i == 0:
            header_row(table.rows[i])
        else:
            table.rows[i]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    # Keep small tables together.
    for row in table.rows[:-1]:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True
    return table


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.2)
section.bottom_margin = Cm(1.0)
section.left_margin = Cm(1.5)
section.right_margin = Cm(1.5)

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(9.0)
normal.font.color.rgb = DARK
normal.paragraph_format.space_after = Pt(1)
normal.paragraph_format.line_spacing = 1.02

# Title block
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0)
r = p.add_run("KAPPA FOUNDATION")
set_font(r, size=21, bold=True, color=NAVY)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(1)
r = p.add_run("Investment Office — Introduction to the Fund Management Board")
set_font(r, size=12.5, bold=True, color=BLUE)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
r = p.add_run("As at 31 December 2022  ·  Prepared for the 22 October 2026 initial submission  ·  All amounts USD")
set_font(r, size=8.5, italic=True, color=GREY)

pPr = p._p.get_or_add_pPr()
pBdr = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
bottom.set(qn("w:val"), "single")
bottom.set(qn("w:sz"), "8")
bottom.set(qn("w:space"), "1")
bottom.set(qn("w:color"), "2E5496")
pBdr.append(bottom)
pPr.append(pBdr)

heading(doc, "WHO WE ARE")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
add_text(
    p,
    "The Trustees have mandated an independent investment team to act as Kappa's investment office. "
    "Our duty: honour every commitment, then pursue the foundation's purchasing-power objective with "
    "disciplined, accountable governance. This page introduces the office; the decisions we recommend "
    "are recorded in the **Initial Decision Record** and the **Committee Memo**.",
    size=9.0,
)

heading(doc, "THE MANDATE AT A GLANCE")
table(
    doc,
    [
        ["Item", "Detail"],
        ["Starting position", "USD 12.0m unrestricted cash, 31 December 2022; no other assets, liabilities or committed donations."],
        ["Commitments", "USD 2.4m scholarship programme due 31 December 2024; USD 3.6m learning centre due 31 December 2027. Fixed, cannot be reduced."],
        ["Objective", "Remaining fund at 31 December 2030 to hold at least USD 8.0m of December-2022 purchasing power — a goal, not a promised return."],
        ["Instruments", "Cash, five default-free case bonds, US common stocks, and US-listed equity ETFs investing primarily in US equities."],
        ["Rules", "No borrowing, short selling or derivatives. USD reporting, tax-exempt, fractional holdings allowed."],
    ],
    widths=[3.0, 15.0],
    font_size=8.4,
)

heading(doc, "OUR INVESTMENT PHILOSOPHY")
bullet(doc, "**Liabilities first** — payment cash flows must not depend on selling risk assets.")
bullet(doc, "**Real goals, real money** — purchasing power is the mission; inflation is the silent liability.")
bullet(doc, "**Risk is mission failure** — we manage the probability and severity of shortfall, not volatility alone.")
bullet(doc, "**Simplicity and low cost** — unless added complexity demonstrably earns its cost.")
bullet(doc, "**Govern before you trade** — decision rules, independent checks and escalation triggers are set in advance.")

heading(doc, "HOW THE OFFICE IS GOVERNED")
table(
    doc,
    [
        ["Body", "Role"],
        ["Board / Trustees", "Approve mission, risk tolerance, permitted assets, benchmarks and major amendments."],
        ["Investment Committee", "Recommends allocation, managers and policy changes; approves limited exceptions."],
        ["Portfolio Manager", "Executes within approved limits; monitors risks and reports."],
        ["Independent Checker", "Re-verifies calculations and challenges assumptions — proposer and checker are never the same person."],
        ["Custodian & Auditor", "Safeguard assets, reconcile cash, verify records."],
    ],
    widths=[4.0, 14.0],
    font_size=8.4,
)

heading(doc, "CAPITAL ARCHITECTURE")
bullet(doc, "**Payment Account** — ring-fenced bond cash flows covering the 2024 and 2027 commitments.")
bullet(doc, "**Liquidity Reserve** — cash for operations, settlement and contingency.")
bullet(doc, "**Growth Account** — diversified US equities pursuing the 2030 purchasing-power goal.")

heading(doc, "HOW TO READ OUR SUBMISSION")
bullet(doc, "**Initial Decision Record (2 pages)** — allocation, universe, benchmark, assumptions, evaluation rule.")
bullet(doc, "**Committee Memo (8 pages + references)** — reasoning, evidence, risks and quantified trade-offs.")
bullet(doc, "**Technical evidence** — reproducible Python, raw data, manifests, hashes and the decision log.")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(4)
p.paragraph_format.space_after = Pt(0)
add_text(
    p,
    "Material AI assistance is disclosed under the course policy; every figure carries human independent checks, recorded separately.",
    size=7.8,
    color=GREY,
)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0)
add_text(p, "Prepared by the Kappa Foundation investment team  ·  Initial decision record due 22 October 2026, 16:00 (GMT+7)", size=7.8, color=GREY)

doc.core_properties.title = "Kappa Foundation — Investment Office Introduction"
doc.core_properties.author = "Kappa investment team"
doc.core_properties.subject = "Fund Management Board introduction, 31 December 2022"

out = Path(r"deliverables\Kappa_Investment_Office_Introduction.docx")
out.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(out))
print(f"Wrote {out}")
