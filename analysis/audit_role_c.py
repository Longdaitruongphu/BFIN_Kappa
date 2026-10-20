"""Independent, read-only numerical audit of the supplied Role C workbook.

This script never edits the workbook or the frozen original policy. It compares
the workbook with saved raw-derived Kappa outputs and checks the workbook's
own regressions. It deliberately does not assert provenance for files absent
from the package.
"""

from __future__ import annotations

import csv
import hashlib
import json
from datetime import datetime
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "reference/role_c/BFIN_Part_C_original.xlsx"
OUT = ROOT / "outputs/role_c_audit"


def csv_by_month(path: Path) -> tuple[list[str], dict[str, dict[str, float]]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        headers = reader.fieldnames or []
        records = {
            row[headers[0]]: {key: float(value) for key, value in row.items() if key != headers[0]}
            for row in reader
        }
    return headers, records


def workbook_months(ws) -> tuple[list[str], dict[str, dict[str, float]]]:
    headers = [str(ws.cell(1, col).value) for col in range(1, 12 if ws.title != "Fama-French Factors" else 6)]
    data: dict[str, dict[str, float]] = {}
    for row in range(2, ws.max_row + 1):
        date = ws.cell(row, 1).value
        if not isinstance(date, datetime):
            continue
        key = date.strftime("%Y-%m")
        data[key] = {
            headers[col - 1]: float(ws.cell(row, col).value)
            for col in range(2, len(headers) + 1)
            if isinstance(ws.cell(row, col).value, (int, float))
        }
    return headers, data


def compare(left: dict[str, dict[str, float]], right: dict[str, dict[str, float]]) -> dict:
    shared_months = sorted(set(left) & set(right))
    shared_tickers = sorted(set.intersection(*(set(d) for d in (left[shared_months[0]], right[shared_months[0]]))))
    details = []
    for ticker in shared_tickers:
        diffs = [(month, left[month][ticker] - right[month][ticker]) for month in shared_months]
        month, difference = max(diffs, key=lambda pair: abs(pair[1]))
        details.append({
            "ticker": ticker,
            "months": len(diffs),
            "max_abs_difference": abs(difference),
            "max_difference_month": month,
            "count_abs_difference_gt_1e_6": sum(abs(value) > 1e-6 for _, value in diffs),
            "count_abs_difference_gt_1e_4": sum(abs(value) > 1e-4 for _, value in diffs),
        })
    return {
        "shared_months": len(shared_months),
        "first_shared_month": shared_months[0],
        "last_shared_month": shared_months[-1],
        "left_only_months": sorted(set(left) - set(right)),
        "right_only_months": sorted(set(right) - set(left)),
        "details": details,
    }


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    book_hash = hashlib.sha256(BOOK.read_bytes()).hexdigest()
    wb = load_workbook(BOOK, data_only=False, read_only=False)
    cached = load_workbook(BOOK, data_only=True, read_only=False)
    project_returns_headers, project_returns = csv_by_month(ROOT / "data/processed/train_monthly_returns.csv")
    project_factor_headers, project_factors = csv_by_month(ROOT / "data/processed/train_factors.csv")
    role_returns_headers, role_returns = workbook_months(wb["Returns_data"])
    role_stocks_headers, role_stocks = workbook_months(wb["Stock Returns"])
    role_factor_headers, role_factors = workbook_months(wb["Fama-French Factors"])
    reg_rows = []
    months = sorted(set(role_stocks) & set(role_factors))
    for row in range(2, 12):
        ws = wb["Multifactor Regression"]
        ticker = ws.cell(row, 1).value
        if not ticker:
            continue
        y = np.array([role_stocks[m][ticker] - role_factors[m]["RF"] for m in months])
        x = np.array([[1, role_factors[m]["Mkt-RF"], role_factors[m]["SMB"], role_factors[m]["HML"]] for m in months])
        beta = np.linalg.lstsq(x, y, rcond=None)[0]
        residual = y - x @ beta
        r2 = 1 - float(residual @ residual) / float(((y - y.mean()) ** 2).sum())
        reported = [float(ws.cell(row, col).value) for col in range(2, 7)]
        recalculated = [*map(float, beta), r2]
        reg_rows.append({
            "ticker": ticker,
            "months": len(months),
            "reported": reported,
            "recalculated": recalculated,
            "max_abs_difference": max(abs(a - b) for a, b in zip(reported, recalculated)),
        })
    checks = wb["C3_return check"]
    checks_cached = cached["C3_return check"]
    spotchecks = []
    for row in range(27, 37):
        spotchecks.append({
            "row": row,
            "ticker": checks.cell(row, 1).value,
            "month": str(checks.cell(row, 2).value),
            "prev_price": checks.cell(row, 3).value,
            "current_price": checks.cell(row, 4).value,
            "calculated_formula": checks.cell(row, 5).value,
            "calculated_cached": checks_cached.cell(row, 5).value,
            "return_link_formula": checks.cell(row, 6).value,
            "return_link_cached": checks_cached.cell(row, 6).value,
            "difference": checks.cell(row, 7).value,
            "status": checks.cell(row, 8).value,
        })
    source = wb["C4_source record"]
    source_refs = []
    for row in range(2, 12):
        rel = source.cell(row, 13).value
        source_refs.append({
            "ticker": source.cell(row, 1).value,
            "reported_file": rel,
            "reported_sha256": source.cell(row, 14).value,
            "exists_in_project": bool(rel and (ROOT / rel).exists()),
            "exists_adjacent_to_supplied_book": bool(rel and (Path("D:/Downloads") / rel).exists()),
        })
    report = {
        "workbook": str(BOOK),
        "sha256": book_hash,
        "project_universe": project_returns_headers[1:],
        "role_c_returns_data_universe": role_returns_headers[1:],
        "role_c_stock_returns_universe": role_stocks_headers[1:],
        "project_factors": project_factor_headers[1:],
        "role_c_factors": role_factor_headers[1:],
        "returns_data_vs_project": compare(role_returns, project_returns),
        "stock_returns_vs_project": compare(role_stocks, project_returns),
        "stock_returns_vs_returns_data": compare(role_stocks, role_returns),
        "factors_vs_project": compare(role_factors, project_factors),
        "regression_rows": reg_rows,
        "spotchecks": spotchecks,
        "source_refs": source_refs,
        "source_summary_sha15": source.cell(15, 8).value,
        "source_summary_sha16": source.cell(16, 8).value,
        "eligibility_checked_by": [wb["C1&2_eligibility"].cell(row, 14).value for row in range(3, 13)],
        "eligibility_evidence_files": [wb["C1&2_eligibility"].cell(row, 13).value for row in range(3, 13)],
        "optimization_nonempty": sum(c.value is not None for row in wb["Optimization"] for c in row),
    }
    output = OUT / "role_c_independent_audit.json"
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(output)
    print("Workbook SHA256:", book_hash)
    print("Return comparison:", json.dumps(report["returns_data_vs_project"], indent=2))
    print("Factor comparison:", json.dumps(report["factors_vs_project"], indent=2))
    print("Regression max difference:", max(row["max_abs_difference"] for row in reg_rows))
    print("Source files present:", sum(ref["exists_in_project"] or ref["exists_adjacent_to_supplied_book"] for ref in source_refs), "/", len(source_refs))


if __name__ == "__main__":
    main()
