"""Read-only numerical and provenance audit of the attached expanded Role C workbook.

Run with the bundled workspace Python (openpyxl, pandas and numpy). This script
never executes workbook macros or treats worksheet claims as instructions.
"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "reference/role_c/BFIN_Part_C_Verified_Expanded_original_20261021.xlsx"
OUT = ROOT / "outputs/role_c_audit/expanded"
OUT.mkdir(parents=True, exist_ok=True)


def number(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return np.nan


def sheet_frame(ws, header_row, first_row, last_row):
    tickers = [str(ws.cell(header_row, c).value) for c in range(2, ws.max_column + 1)]
    rows = []
    for r in range(first_row, last_row + 1):
        date = str(ws.cell(r, 1).value)[:7]
        rows.append([date] + [number(ws.cell(r, c).value) for c in range(2, ws.max_column + 1)])
    frame = pd.DataFrame(rows, columns=["month"] + tickers).set_index("month")
    return frame


def stats_monthly(monthly):
    monthly = np.asarray(monthly, dtype=float)
    wealth = np.cumprod(1 + monthly)
    peak = np.maximum.accumulate(np.r_[1, wealth])[1:]
    return {"cagr": float(wealth[-1] ** (12 / len(monthly)) - 1),
            "vol": float(np.std(monthly, ddof=1) * np.sqrt(12)),
            "max_drawdown": float(np.min(wealth / peak - 1)),
            "terminal_multiple": float(wealth[-1])}


def main():
    wb = load_workbook(SOURCE, read_only=False, data_only=False)
    values = load_workbook(SOURCE, read_only=False, data_only=True)
    source_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    formula_count = cache_count = 0
    for ws in wb:
        dv = values[ws.title]
        for row in ws:
            for cell in row:
                if cell.data_type == "f":
                    formula_count += 1
                    if dv[cell.coordinate].value is not None:
                        cache_count += 1

    prices = sheet_frame(wb["Raw_Adjusted_Prices"], 4, 5, 137)
    prices.to_csv(OUT / "expanded_raw_prices.csv")
    ret = prices.pct_change().iloc[1:]
    ret.to_csv(OUT / "expanded_recomputed_returns.csv")
    ff = sheet_frame(wb["Fama-French Factors"], 1, 2, 97)
    ff.to_csv(OUT / "expanded_factors.csv")
    can_prices = pd.read_csv(ROOT / "data/processed/train_month_end_adjusted_prices.csv", index_col="month")
    can_ff = pd.read_csv(ROOT / "data/processed/train_factors.csv", index_col="month")
    overlap = sorted(set(prices.columns).intersection(can_prices.columns))
    price_checks = []
    for t in overlap:
        d = prices.loc[can_prices.index, t] - can_prices[t]
        retd = ret.loc[ret.index.intersection(can_prices.index), t] - can_prices[t].pct_change().loc[ret.index.intersection(can_prices.index)]
        price_checks.append({"ticker": t, "max_abs_price_difference": float(d.abs().max()),
                             "mean_abs_price_difference": float(d.abs().mean()),
                             "price_disagreements_over_1cent": int((d.abs() > .01).sum()),
                             "max_abs_monthly_return_difference": float(retd.abs().max()),
                             "dec2014_expanded": float(prices.loc["2014-12", t]),
                             "dec2014_canonical": float(can_prices.loc["2014-12", t])})
    pd.DataFrame(price_checks).to_csv(OUT / "price_source_comparison.csv", index=False)
    ff_diff = (ff.loc[can_ff.index] - can_ff[ff.columns]).abs()
    ff_checks = {c: {"max_abs_difference": float(ff_diff[c].max()),
                     "nonzero_cells": int((ff_diff[c] > 1e-10).sum())} for c in ff.columns}
    can_eval_ret = pd.read_csv(ROOT / "data/processed/eval_monthly_returns.csv", index_col="month")
    oos_checks = []
    for t in overlap:
        difference = ret.loc[can_eval_ret.index, t] - can_eval_ret[t]
        worst = difference.abs().idxmax()
        oos_checks.append({"ticker": t, "max_abs_monthly_return_difference": float(difference.abs().max()),
                           "worst_month": worst,
                           "expanded_return_worst_month": float(ret.loc[worst, t]),
                           "saved_yahoo_return_worst_month": float(can_eval_ret.loc[worst, t]),
                           "expanded_2023_25_cagr": float(np.prod(1 + ret.loc[can_eval_ret.index, t]) ** (1 / 3) - 1),
                           "saved_yahoo_2023_25_cagr": float(np.prod(1 + can_eval_ret[t]) ** (1 / 3) - 1)})
    pd.DataFrame(oos_checks).to_csv(OUT / "evaluation_source_comparison.csv", index=False)

    opt = wb["Optimization"]
    weights = pd.DataFrame({"ticker": [opt.cell(r, 1).value for r in range(5, 29)],
                            "conservative": [number(opt.cell(r, 4).value) for r in range(5, 29)],
                            "baseline": [number(opt.cell(r, 5).value) for r in range(5, 29)],
                            "aggressive": [number(opt.cell(r, 6).value) for r in range(5, 29)]})
    weights.to_csv(OUT / "expanded_weights.csv", index=False)
    reg = wb["Multifactor Regression"]
    pd.DataFrame([{"ticker": reg.cell(r, 1).value,
                   "reported_alpha_monthly": number(reg.cell(r, 4).value),
                   "reported_alpha_t": number(reg.cell(r, 7).value),
                   "reported_alpha_p": number(reg.cell(r, 8).value),
                   "reported_beta_mkt": number(reg.cell(r, 11).value)}
                  for r in range(5, 44)]).to_csv(OUT / "expanded_regression_claims.csv", index=False)
    portfolio_checks = {}
    for name, excel_row in [("conservative", 34), ("baseline", 35), ("aggressive", 36)]:
        w = weights.set_index("ticker")[name]
        rtrain = ret.loc["2015-01":"2022-12", w.index].to_numpy() @ w.to_numpy()
        roos = ret.loc["2023-01":"2025-12", w.index].to_numpy() @ w.to_numpy()
        fixed_train = stats_monthly(rtrain)
        fixed_oos = stats_monthly(roos)
        # Rebalanced monthly (the convention most likely underlying weighted return tables).
        claimed = {"train_cagr": number(opt.cell(excel_row, 3).value),
                   "train_vol": number(opt.cell(excel_row, 4).value),
                   "train_max_drawdown": number(opt.cell(excel_row, 6).value),
                   "oos_cagr": number(opt.cell(excel_row, 10).value),
                   "oos_max_drawdown": number(opt.cell(excel_row, 11).value)}
        portfolio_checks[name] = {"weight_sum": float(w.sum()), "number_positive": int((w > 0).sum()),
                                  "claim": claimed, "monthly_rebalanced_train": fixed_train,
                                  "monthly_rebalanced_oos": fixed_oos,
                                  "claim_minus_recomputed_train_cagr": claimed["train_cagr"] - fixed_train["cagr"],
                                  "claim_minus_recomputed_oos_cagr": claimed["oos_cagr"] - fixed_oos["cagr"]}

    # Workbook labels a 10-by-10 lower triangle sample and upper triangle SIM.
    names = [str(opt.cell(54, c).value) for c in range(2, 12)]
    s = ret.loc["2015-01":"2022-12", names].cov() * 12
    cchecks = []
    for i in range(10):
        for j in range(i + 1):
            saved = number(opt.cell(55 + i, 2 + j).value)
            cchecks.append({"asset_i": names[i], "asset_j": names[j], "saved_sample": saved,
                            "recomputed_sample_from_C_prices": float(s.iloc[i, j]),
                            "abs_difference": abs(saved - s.iloc[i, j])})
    pd.DataFrame(cchecks).to_csv(OUT / "covariance_sample_comparison.csv", index=False)

    audit = {
        "source": str(SOURCE), "sha256": source_hash,
        "formula_count": formula_count, "formulas_with_cached_result": cache_count,
        "sheets": {x.title: {"rows": x.max_row, "columns": x.max_column} for x in wb},
        "expanded_assets": list(prices.columns), "asset_count": len(prices.columns),
        "months_prices": list(prices.index[[0, -1]]), "monthly_return_rows": len(ret),
        "price_overlap_with_saved_canonical": price_checks,
        "oos_return_overlap_with_saved_canonical": oos_checks,
        "ff3_against_saved_canonical": ff_checks,
        "portfolios": portfolio_checks,
        "covariance_sample_max_abs_error": max(c["abs_difference"] for c in cchecks),
        "covariance_sample_error_cells_over_1e-8": int(sum(c["abs_difference"] > 1e-8 for c in cchecks)),
        "cell_evidence": {
            "case_conflict_spending": "0_Audit_&_Endowment_IPS!B14:B18",
            "lambda_3_5": "Optimization!E4:E29; Optimization!B35",
            "hardcoded_weights": "Optimization!D5:F28",
            "hardcoded_claimed_performance": "Optimization!C34:G38; J34:K38",
            "unexecuted_source_check": "C4_source record!N3:N41",
            "generic_source_links": "Raw_Adjusted_Prices!B2:AN3",
            "hardcoded_factors": "Fama-French Factors!B2:E97",
        },
    }
    (OUT / "expanded_workbook_audit.json").write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps({k: audit[k] for k in ["sha256", "formula_count", "formulas_with_cached_result", "asset_count", "ff3_against_saved_canonical", "covariance_sample_max_abs_error", "covariance_sample_error_cells_over_1e-8"]}, indent=2))
    print("Saved:", OUT)


if __name__ == "__main__":
    main()
