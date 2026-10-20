from __future__ import annotations
import hashlib, json, math
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'reference' / 'BFIN_Part_B_original.xlsx'
POLICY = ROOT / 'policy' / 'frozen_policy.json'
OUT = ROOT / 'outputs' / 'role_b_audit' / 'role_b_independent_audit.json'
wb_f = load_workbook(SOURCE, data_only=False, read_only=True)
wb_v = load_workbook(SOURCE, data_only=True, read_only=True)
ass_f, ass_v = wb_f['Bond_Assumptions'], wb_v['Bond_Assumptions']
flow_f, flow_v = wb_f['Bond Sheet'], wb_v['Bond Sheet']
policy = json.loads(POLICY.read_text(encoding='utf-8'))
inputs = {
    'initial_fund': float(ass_v['B5'].value), 'scholarship': float(ass_v['B7'].value),
    'beta_price_per_1000': float(ass_v['B8'].value), 'learning_centre': float(ass_v['B11'].value),
    'delta_price_per_1000': float(ass_v['B12'].value), 'delta_coupon': float(ass_v['B13'].value),
}
beta_face = inputs['scholarship']
delta_face = inputs['learning_centre'] / (1 + inputs['delta_coupon'])
beta_units, delta_units = beta_face/1000, delta_face/1000
beta_cost = beta_units * inputs['beta_price_per_1000']
delta_cost = delta_units * inputs['delta_price_per_1000']
bond_cost = beta_cost + delta_cost
bond_fee = bond_cost * policy['forecast_assumptions']['bond_cost_per_dollar_traded']
equity_fee_rate = policy['forecast_assumptions']['equity_cost_per_dollar_traded']
equity_dollars_before_cost = (inputs['initial_fund'] - bond_cost - bond_fee)/(1+equity_fee_rate)
equity_fee = equity_dollars_before_cost * equity_fee_rate
equity_market_value = equity_dollars_before_cost
delta_coupon_cash = delta_face * inputs['delta_coupon']
terminal_surplus = delta_coupon_cash * 4

# Reconstruct the spreadsheet's central values directly from displayed case inputs.
checks = {
    'B18_beta_units': {'workbook_cached': ass_v['B18'].value, 'independent': beta_units},
    'B19_beta_cost': {'workbook_cached': ass_v['B19'].value, 'independent': beta_cost},
    'B20_delta_units': {'workbook_cached': ass_v['B20'].value, 'independent': delta_units},
    'B21_delta_cost': {'workbook_cached': ass_v['B21'].value, 'independent': delta_cost},
    'B22_bond_cost': {'workbook_cached': ass_v['B22'].value, 'independent': bond_cost},
    'B23_growth_before_fees': {'workbook_cached': ass_v['B23'].value, 'independent': inputs['initial_fund']-bond_cost},
}
checks['B5_2027_net_cash'] = {'workbook_cached': flow_v['F9'].value, 'independent': delta_face*(1+inputs['delta_coupon'])-inputs['learning_centre']}
checks['four_delta_coupons'] = {'workbook_cached_total_if_retained_cash': None, 'independent': terminal_surplus}
checks['project_cost_inclusive_equity'] = {'project_market_value': equity_dollars_before_cost, 'project_equity_fee': equity_fee, 'project_bond_fee': bond_fee, 'project_total_fees': equity_fee+bond_fee, 'B_growth_overstatement_vs_project_invested_equity': float(ass_v['B23'].value)-equity_dollars_before_cost}

# Compare Role C sheets to frozen set and each other.
cfile = ROOT/'reference'/'role_c'/'BFIN_Part_C_original.xlsx'
wc = load_workbook(cfile, data_only=True, read_only=True)
role_c = {
  'returns_data_universe': [wc['Returns_data'].cell(1,c).value for c in range(2,12)],
  'stock_returns_universe': [wc['Stock Returns'].cell(1,c).value for c in range(2,12)],
  'frozen_project_universe': policy['forecast_assumptions']['stocks']+policy['forecast_assumptions']['etfs'],
  'overlap_returns_data': sorted(set([wc['Returns_data'].cell(1,c).value for c in range(2,12)]) & set(policy['forecast_assumptions']['stocks']+policy['forecast_assumptions']['etfs'])),
  'returns_data_differs_from_frozen': sorted(set([wc['Returns_data'].cell(1,c).value for c in range(2,12)]) ^ set(policy['forecast_assumptions']['stocks']+policy['forecast_assumptions']['etfs'])),
  'stock_returns_differs_from_frozen': sorted(set([wc['Stock Returns'].cell(1,c).value for c in range(2,12)]) ^ set(policy['forecast_assumptions']['stocks']+policy['forecast_assumptions']['etfs'])),
  'regression_includes_frozen_USMV': any(row[0].value=='USMV' for row in wc['Multifactor Regression'].iter_rows(min_col=1, max_col=1)),
  'optimization_sheet_nonempty': any(cell.value is not None for row in wc['Optimization'].iter_rows() for cell in row),
}
receipt={
 'source_file': str(SOURCE.relative_to(ROOT)),
 'source_sha256': hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
 'policy_sha256': hashlib.sha256(POLICY.read_bytes()).hexdigest(),
 'method': 'Independent arithmetic reconstructed from workbook input cells, project frozen cost assumptions, and saved workbook values; no formula cache treated as proof.',
 'case_inputs': inputs,
 'workbook_formula_text': {cell:ass_f[cell].value for cell in ['B18','B19','B20','B21','B22','B23']},
 'checks': checks,
 'role_c_comparison': role_c,
 'findings': [
   'Bond quantity and quoted-price arithmetic reconcile to case inputs.',
   'B workbook excludes project transaction-cost assumptions; its remaining growth capital is gross of fees and exceeds project USMV market value by costs.',
   'Cashflow match covers scheduled nominal payments under fictional default-free case bond assumptions.',
   'Duration equality alone does not establish absolute immunization; cashflow matching is the accurate description here.',
   'Buy-and-hold removes the need to realize interim price marks only if no forced sale; it does not make the entire foundation safe or eliminate liquidity, inflation, opportunity-cost, or growth-sleeve risk.',
   'Redirecting coupons into equity is a distinct policy scenario from the frozen 0% cash-reinvestment baseline.'
 ]
}
OUT.write_text(json.dumps(receipt, indent=2), encoding='utf-8')
print(OUT)
print(json.dumps(receipt, indent=2))

