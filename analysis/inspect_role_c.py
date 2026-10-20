"""Read-only inspection of the supplied Role C workbook."""

from __future__ import annotations

import argparse
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser()
    parser.add_argument("--sheet")
    parser.add_argument("--first", type=int, default=1)
    parser.add_argument("--last", type=int, default=80)
    args = parser.parse_args()

    source = Path("reference/role_c/BFIN_Part_C_original.xlsx")
    formulas = load_workbook(source, data_only=False, read_only=False)
    values = load_workbook(source, data_only=True, read_only=False)
    for ws in formulas.worksheets:
        nonempty = [(c.coordinate, c.value) for row in ws for c in row if c.value is not None]
        formula_cells = [(coord, value) for coord, value in nonempty if isinstance(value, str) and value.startswith("=")]
        error_cells = [(c.coordinate, c.value) for row in values[ws.title] for c in row if c.data_type == "e"]
        types = Counter(c.data_type for row in ws for c in row if c.value is not None)
        print(f"SHEET {ws.title!r}: nonempty={len(nonempty)}, formula={len(formula_cells)}, cached_errors={len(error_cells)}, types={dict(types)}")
        if error_cells:
            print("  ERRORS", error_cells[:20])
        if args.sheet and ws.title != args.sheet:
            continue
        for row in ws.iter_rows(min_row=args.first, max_row=min(args.last, ws.max_row)):
            filled = [f"{c.coordinate}={str(c.value)[:150]}" for c in row if c.value is not None]
            if filled:
                print("  ", " | ".join(filled))


if __name__ == "__main__":
    main()
