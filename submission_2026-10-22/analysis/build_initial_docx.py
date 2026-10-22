"""Build an editable, two-page Initial Decision Record from the saved model outputs.

Run from the dated submission root after ``python analysis/run_initial.py``.
The PDF remains the print master; this document gives the team an editable copy
with the same decision, amounts, assumptions, and evaluation rule.
"""
from __future__ import annotations

import csv
import io
import json
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables/Kappa_Initial_Decision_Record_22Oct2026.docx"
DATA = json.loads((ROOT / "outputs/initial_decision_rebuild.json").read_text(encoding="utf-8"))
ALT = json.loads((ROOT / "outputs/sample_utility_robustness.json").read_text(encoding="utf-8"))
with (ROOT / "outputs/risk_preference.csv").open(encoding="utf-8", newline="") as f:
    RISK = list(csv.DictReader(f))
with (ROOT / "outputs/max_utility_preference.csv").open(encoding="utf-8", newline="") as f:
    MAX_UTILITY = list(csv.DictReader(f))

NAVY = RGBColor(0, 42, 84)
BLUE = RGBColor(0, 114, 206)
TEAL = RGBColor(0, 137, 153)
INK = RGBColor(26, 32, 44)
MUTED = RGBColor(90, 107, 124)
PALE = "EBF0F5"
LINE = "D9E1E8"


def set_cell_fill(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=65, start=80, bottom=65, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    m = tc_pr.first_child_found_in("w:tcMar")
    if m is None:
        m = OxmlElement("w:tcMar")
        tc_pr.append(m)
    for tag, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = m.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            m.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def no_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "nil")


def keep_rows(table):
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        el = OxmlElement("w:cantSplit")
        tr_pr.append(el)


def style_run(run, size=8.35, bold=False, color=INK, italic=False, super_=False):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    if super_:
        run.font.superscript = True
    return run


def set_para(p, before=0, after=0, line=1.08, align=None):
    f = p.paragraph_format
    f.space_before = Pt(before)
    f.space_after = Pt(after)
    f.line_spacing = line
    if align is not None:
        p.alignment = align
    return p


def add_text(container, text, *, size=8.35, bold=False, color=INK,
             italic=False, before=0, after=3, line=1.1):
    p = set_para(container.add_paragraph(), before, after, line)
    style_run(p.add_run(text), size, bold, color, italic)
    return p


def add_heading(container, text, *, before=7, after=3):
    p = set_para(container.add_paragraph(), before, after, 1.0)
    style_run(p.add_run(text.upper()), 9.0, True, NAVY)
    return p


def add_source(container, text):
    p = set_para(container.add_paragraph(), 0, 4, 1.0)
    style_run(p.add_run(text), 7.5, False, MUTED, True)
    return p


def add_header(doc, title, subtitle):
    p = set_para(doc.add_paragraph(), after=3, line=1)
    style_run(p.add_run("KAPPA FOUNDATION  /  INVESTMENT COMMITTEE"), 7.9, True, TEAL)
    p = set_para(doc.add_paragraph(style="Title"), after=2, line=1)
    style_run(p.add_run(title), 18.5, True, NAVY)
    p = set_para(doc.add_paragraph(), after=6, line=1)
    style_run(p.add_run(subtitle), 8.1, False, MUTED)


def add_table(container, headers, rows, widths, align_right=()):
    t = container.add_table(rows=1, cols=len(headers))
    t.autofit = False
    no_borders(t)
    for i, (h, w) in enumerate(zip(headers, widths)):
        c = t.rows[0].cells[i]
        c.width = Cm(w)
        set_cell_fill(c, "002A54")
        set_cell_margins(c)
        p = set_para(c.paragraphs[0], line=1)
        if i in align_right:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        style_run(p.add_run(h), 7.6, True, RGBColor(255, 255, 255))
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, (v, w) in enumerate(zip(row, widths)):
            c = cells[i]
            c.width = Cm(w)
            set_cell_margins(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri % 2:
                set_cell_fill(c, PALE)
            p = set_para(c.paragraphs[0], line=1.02)
            if i in align_right:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            val = str(v)
            marker = None
            for sup in "¹²³⁴⁵":
                if val.endswith(sup):
                    val, marker = val[:-1], sup
                    break
            style_run(p.add_run(val), 7.7, i == 0, INK)
            if marker:
                style_run(p.add_run(marker), 5.5, False, INK, super_=True)
    keep_rows(t)
    return t


def add_fact_table(container, rows, widths=(2.6, 2.3)):
    return add_table(container, ("FACT", "VALUE"), rows, widths)


def capital_chart():
    vals = [DATA["beta_initial_value"], DATA["delta_initial_value"], DATA["growth_initial_value"],
            DATA["bond_fee"] + DATA["growth_fee"]]
    names = ["Beta 18.46%", "Delta 28.58%", "USMV 52.89%", "Costs 0.08%"]
    colors = ["#002A54", "#0072CE", "#008999", "#CBD6DF"]
    im = Image.new("RGBA", (1760, 95), (255, 255, 255, 0))
    draw = ImageDraw.Draw(im)
    left = 0
    for v, col in zip(vals, colors):
        width = round(1760 * v / 12_000_000)
        draw.rectangle((left, 18, min(1759, left + width), 73), fill=col)
        left += width
    stream = io.BytesIO()
    im.save(stream, format="PNG")
    stream.seek(0)
    return stream, names, colors


def utility_chart():
    a = [int(x["risk_aversion_A"]) for x in RISK]
    u = [100 * float(x["USMV_utility"]) for x in RISK]
    v = [100 * float(x["VTI_utility"]) for x in RISK]
    m = [100 * float(x["utility"]) for x in MAX_UTILITY]
    im = Image.new("RGB", (1760, 320), "white")
    draw = ImageDraw.Draw(im)
    font_file = Path("C:/Windows/Fonts/arial.ttf")
    font = ImageFont.truetype(str(font_file), 25) if font_file.exists() else ImageFont.load_default()
    bold_file = Path("C:/Windows/Fonts/arialbd.ttf")
    bold = ImageFont.truetype(str(bold_file), 25) if bold_file.exists() else font
    x0, x1, y0, y1 = 95, 1700, 28, 245
    low, high = -3.0, 7.0
    xy = lambda aa, uu: (round(x0 + (aa - 2) / 6 * (x1 - x0)),
                         round(y1 - (uu - low) / (high - low) * (y1 - y0)))
    for tick in (-2, 0, 2, 4, 6):
        yy = xy(2, tick)[1]
        draw.line((x0, yy, x1, yy), fill="#E7EDF2", width=2)
        draw.text((24, yy - 13), str(tick), fill="#5A6B7C", font=font)
    for aa in a:
        xx = xy(aa, 0)[0]
        draw.text((xx - 9, y1 + 9), str(aa), fill="#5A6B7C", font=font)
    xx = xy(4, 0)[0]
    draw.line((xx, y0, xx, y1), fill="#CCD8E1", width=3)
    for scores, col, dashed, square in ((u, "#002A54", False, False),
                                         (v, "#008999", False, False),
                                         (m, "#0072CE", True, True)):
        pts = [xy(aa, score) for aa, score in zip(a, scores)]
        if dashed:
            for p0, p1 in zip(pts[:-1], pts[1:]):
                dx, dy = p1[0] - p0[0], p1[1] - p0[1]
                length = max(1, (dx * dx + dy * dy) ** .5)
                n = max(1, int(length / 20))
                for j in range(0, n, 2):
                    t0, t1 = j / n, min((j + 1) / n, 1)
                    draw.line((round(p0[0] + dx*t0), round(p0[1] + dy*t0),
                               round(p0[0] + dx*t1), round(p0[1] + dy*t1)), fill=col, width=5)
        else:
            draw.line(pts, fill=col, width=5, joint="curve")
        for px, py in pts:
            if square:
                draw.rectangle((px - 8, py - 8, px + 8, py + 8), fill=col)
            else:
                draw.ellipse((px - 7, py - 7, px + 7, py + 7), fill=col)
    draw.text((x0 + 25, 2), "Utility (%)", fill="#5A6B7C", font=font)
    draw.line((x0, 285, x0 + 45, 285), fill="#002A54", width=5)
    draw.text((x0 + 53, 272), "USMV 100%", fill="#002A54", font=font)
    draw.line((x0 + 250, 285, x0 + 295, 285), fill="#008999", width=5)
    draw.text((x0 + 303, 272), "VTI 100%", fill="#008999", font=font)
    draw.line((x0 + 510, 285, x0 + 555, 285), fill="#0072CE", width=5)
    draw.text((x0 + 563, 272), "Max utility, re-optimised", fill="#0072CE", font=font)
    draw.text((x1 - 205, 282), "Risk aversion A", fill="#5A6B7C", font=font)
    stream = io.BytesIO()
    im.save(stream, format="PNG")
    stream.seek(0)
    return stream


def make():
    assert DATA["training_months"] == 96
    assert abs(sum(DATA[k] for k in ("beta_initial_value", "delta_initial_value",
                                     "bond_fee", "growth_initial_value", "growth_fee")) - 12e6) < 1e-6
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(1.55), Cm(1.35)
    sec.left_margin, sec.right_margin = Cm(1.55), Cm(1.55)
    sec.header_distance = Cm(.5)
    sec.footer_distance = Cm(.55)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(8.35)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(3)
    title_style = doc.styles["Title"]
    title_style.font.name = "Arial"
    title_style.font.color.rgb = NAVY
    title_style.font.size = Pt(18.5)
    title_style.paragraph_format.space_after = Pt(2)
    title_pr = title_style._element.get_or_add_pPr()
    title_border = title_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_pr.remove(title_border)

    footer = sec.footer.paragraphs[0]
    set_para(footer, align=WD_ALIGN_PARAGRAPH.RIGHT)
    style_run(footer.add_run("Prepared 21 October 2026  |  Case information cutoff 31 December 2022"),
              7.1, False, MUTED)

    add_header(doc, "Initial decision record", "USD  |  Case date 31 December 2022  |  Initial submission 22 October 2026")
    add_heading(doc, "Investment objective and policy", before=0)
    add_text(doc, "Meet the fixed education commitments first and protect long-term purchasing power. Invest surplus capital within the fund's capacity for risk. The end-2030 goal is USD 8m in December 2022 purchasing power; it is a goal, not a promised return. Changes require evidence about obligations, resources and assets.", after=4)

    add_heading(doc, "Initial capital allocation", before=4)
    stream, names, cols = capital_chart()
    doc.add_picture(stream, width=Cm(17.7), height=Cm(.42))
    legend = set_para(doc.add_paragraph(), after=3)
    for name, col in zip(names, cols):
        marker = legend.add_run("● ")
        style_run(marker, 7.3, False, RGBColor.from_string(col[1:]))
        style_run(legend.add_run(name + "     "), 7.3, False, MUTED)
    add_table(doc, ("HOLDING / USE", "UNITS / BASIS", "INITIAL USD"), [
        ("Beta | 2024", "2,400", "2,214,672.00"),
        ("Delta | 2027", "3,478.26087", "3,429,808.70"),
        ("USMV", "100% of sleeve", "6,346,350.71"),
        ("Trading costs", "5 / 10bp", "9,168.59"),
        ("Total", "", "12,000,000.00"),
    ], (6.4, 4.9, 6.6), align_right=(2,))
    add_source(doc, "Source: case Exhibit 1; Kappa model; USD, fees separate; percentages rounded.")
    add_text(doc, "Beta principal pays $2.4m on 31 Dec 2024. Delta principal and final coupon pay $3.6m on 31 Dec 2027. Four earlier Delta coupons total $486,957 and remain cash at 0%; payment coverage needs no equity sale or reinvestment return.", after=2)

    add_heading(doc, "Growth portfolio decision", before=4)
    add_text(doc, "Recommend 100% USMV, the US minimum-volatility ETF. Use 100% VTI, a broad-US-market ETF, as the frozen investable benchmark and credible rejected alternative.", after=3)
    usmv = DATA["forecast_comparison"]["USMV"]
    vti = DATA["forecast_comparison"]["VTI"]
    add_table(doc, ("ANNUAL MEASURE", "USMV", "VTI", "MAX UTILITY A=4"), [
        ("Expected return¹", f"{usmv['forecast_return_annual']:.2%}", f"{vti['forecast_return_annual']:.2%}", f"{ALT['expected_return']:.2%}"),
        ("Volatility²", f"{usmv['historical_volatility_annual']:.2%}", f"{vti['historical_volatility_annual']:.2%}", f"{ALT['volatility']:.2%}"),
        ("Utility³ | A=4", f"{usmv['utility_A4']:.2%}", f"{vti['utility_A4']:.2%}", f"{ALT['utility_A4']:.2%}"),
    ], (6.0, 3.9, 3.9, 4.1), align_right=(1, 2, 3))
    add_source(doc, "Source: Yahoo Finance, Kenneth French; Kappa model; 2015-22 monthly returns, annualised.")
    w = ALT["weights_96m"]
    weights1 = " | ".join(f"{t} {w[t]:.2%}" for t in ("MSFT", "JPM", "JNJ", "PG"))
    weights2 = " | ".join(f"{t} {w[t]:.2%}" for t in ("XOM", "NEE", "USMV", "VTI")) + " | AAPL/CAT 0.00%"
    p = add_text(doc, "Full-universe maximum utility at A=4 (96 months, sample covariance):", after=1)
    for run in p.runs: run.bold = True
    add_text(doc, weights1 + "\n" + weights2, after=2)
    gap_pp = 100 * (ALT["utility_A4"] - DATA["forecast_comparison"]["USMV"]["utility_A4"])
    add_text(doc, f"This is the highest-scoring long-only mix of the 10 candidates for the stated inputs: CAPM expected returns, sample covariance, weights >=0 and total weight=100%, with no extra caps. It estimates 7.19% return, 12.70% volatility and 3.96% utility, {gap_pp:.2f} percentage point above USMV. The 96-to-60-month check shifts 34.8% of weights one-way (about $2.21m on this sleeve). We choose USMV for simpler policy implementation, not because it is the model maximum.", after=2)

    detail = doc.add_table(rows=1, cols=2)
    detail.autofit = False
    no_borders(detail)
    left, right = detail.rows[0].cells
    left.width, right.width = Cm(8.9), Cm(8.9)
    for cell in (left, right):
        set_cell_margins(cell, 90, 95, 70, 95)
        set_cell_fill(cell, PALE)
    p = set_para(left.paragraphs[0], after=2)
    style_run(p.add_run("MANDATE & ASSUMPTIONS"), 7.8, True, NAVY)
    for s in ("$12m initial; USD 8m real at 2030; CPI 2.5% base, 3.0% education-cost stress.",
              "Forecast RF 4%; market premium 4.5%; forecast alpha 0; provisional A=4.",
              "Coupon cash 0%; 5bp bond / 10bp equity entry cost; no extra position caps."):
        p = set_para(left.add_paragraph(), after=2, line=1.02)
        style_run(p.add_run(s), 7.45)
    p = set_para(right.paragraphs[0], after=2)
    style_run(p.add_run("CANDIDATE UNIVERSE"), 7.8, True, NAVY)
    for s in ("Technology AAPL, MSFT; financials JPM; health care JNJ; staples PG.",
              "Energy XOM; industrials CAT; utilities NEE; equity ETFs USMV and VTI.",
              "Eight stocks in seven sectors; two ETF mandates. Candidates need not be held."):
        p = set_para(right.add_paragraph(), after=2, line=1.02)
        style_run(p.add_run(s), 7.45)
    add_source(doc, "Source: original case, proposed IPS and dated eligibility notes; instruments are long-only, without borrowing, shorting or derivatives.")

    doc.add_page_break()
    add_header(doc, "Policy, risk and implementation", "Growth-sleeve signals  |  Funding and 2030 goal tracked separately  |  Proposed policy")
    add_heading(doc, "Risk preference and real goal", before=0)
    doc.add_picture(utility_chart(), width=Cm(17.7), height=Cm(3.15))
    add_source(doc, "Source: Kappa analysis from saved Yahoo/French inputs; utility = expected return - 0.5 x A x annual variance.")
    add_text(doc, "USMV and VTI stay at 100%. The dashed line re-optimises all 10 candidates at each A and shows the highest estimated score for that risk preference. At A=4 it scores 3.96%; page 1 lists its holdings. Of the fixed ETFs, VTI leads at A=2; USMV leads at A=4, 6 and 8 (crossover A=2.58). A=4 is provisional.", after=2)
    add_table(doc, ("2030 GOAL", "CPI BASE", "EDUCATION STRESS"), [
        ("Inflation", "2.50%", "3.00%"),
        ("Nominal USD", "$9.747m", "$10.134m"),
        ("CAGR⁴ hurdle", "4.84%", "5.37%"),
    ], (6.1, 5.3, 6.5), align_right=(1, 2))
    add_source(doc, "Source: Kappa funding model; 2022-30; four pre-maturity coupons retained at 0% cash interest.")

    add_heading(doc, "Downside oversight", before=4)
    add_text(doc, "Primary signal: flow-adjusted growth-sleeve total-return drawdown from its running high, including cash; retain the prior peak after changes. Check whole-fund stress and dated-payment coverage separately. These are review triggers, not guaranteed loss limits.", after=2)
    add_table(doc, ("SIGNAL⁵", "REQUIRED ACTION", "DEADLINE"), [
        ("10%", "Validate data, payment coverage and stress", "2 business days"),
        ("15%", "Investment Committee decision", "5 business days"),
        ("20%", "Notify trustees; pause risk increases", "Same / next day"),
        ("Payment gap", "Escalate; restore coverage", "Immediate"),
    ], (3.1, 10.2, 4.6))
    add_source(doc, "Source: proposed IPS, 21 Oct 2026; signal is month-end growth-sleeve drawdown, except a payment gap at any time.")
    add_text(doc, "No automatic sale. At 15%, record a hold, authorised rebalance or prospective policy proposal. A crisis may defer discretionary action only with written authority, expiry and review; due-payment coverage cannot fall.", after=2)

    add_heading(doc, "Annual goal protection and policy response", before=4)
    add_text(doc, "At each annual review, first confirm remaining bond and cash obligations. Test the 2030 goal of $9.747m nominal at CPI 2.5% ($10.134m at HEPI 3% stress) after a 30% fall in growth assets. If stressed resources still cover the goal, trustees may approve locking only the minimum needed in a permitted, actually priced asset due by 2030; keep excess invested. This is a prospective trustee decision, not automatic profit-taking. If a gap remains, report it; do not raise risk to catch up. For relative underperformance, diagnose exposures, costs and implementation before proposing a change.", after=2)
    add_heading(doc, "Frozen test and future rebalancing", before=3)
    add_text(doc, "For the test, freeze USMV 100% and VTI 100% of the growth sleeve at 31 Dec 2022. Compare monthly total returns in Jan 2023-Dec 2025. Reset to starting weights after each December return, without updating forecasts. Since each portfolio holds one ETF, this rule normally creates no trade; charge 10bp only when a trade occurs. Later returns test the frozen choice; they do not set the 2022 weights. Future changes require a dated committee review under the IPS.", after=2)
    add_text(doc, "Estimate Jan 2015-Dec 2022 (96 common months); a 60-month sample check informs the initial recommendation. For the final memo, compare sample with single-index covariance. CAPM forecasts set alpha to zero; FF3 explains historical exposures. No 2023-25 observation enters the 2022 weights.", after=3)

    mini = doc.add_table(rows=1, cols=2)
    mini.autofit = False
    no_borders(mini)
    c1, c2 = mini.rows[0].cells
    c1.width, c2.width = Cm(8.9), Cm(8.9)
    for c in (c1, c2):
        set_cell_margins(c, 85, 95, 65, 95)
        set_cell_fill(c, PALE)
    p = set_para(c1.paragraphs[0], after=2)
    style_run(p.add_run("FIDUCIARY SCOPE"), 7.7, True, NAVY)
    p = set_para(c1.add_paragraph(), line=1.02)
    style_run(p.add_run("Nominal payments are covered under the case's default-free bond assumption. Neither ETF, the modelled return nor the 2030 real goal is guaranteed; trustees must authorise the proposal."), 7.2)
    p = set_para(c2.paragraphs[0], after=2)
    style_run(p.add_run("QUANTITATIVE AUDIT"), 7.7, True, NAVY)
    p = set_para(c2.add_paragraph(), line=1.02)
    style_run(p.add_run("Run analysis/run_initial.py from saved raw data. Inspect outputs/max_utility_A4_weights.csv for tickers/weights and max_utility_preference.csv for the A grid; see DATA_RECORD.md and retrieval_manifest.jsonl for provenance. This is a retrospective case."), 7.2)

    add_heading(doc, "Mini glossary", before=5, after=2)
    defs = [
        ("¹ CAPM return", "Annual arithmetic forecast: risk-free rate + beta x market premium; forecast alpha=0."),
        ("² Volatility", "Monthly simple-return standard deviation x sqrt(12); a risk estimate, not a loss cap."),
        ("³ Utility (A)", "E[R] - 0.5 x A x annual variance; a larger A penalises risk more."),
        ("⁴ CAGR", "The constant annual compound rate; the goal hurdle includes retained coupon cash."),
        ("⁵ Drawdown", "Decline in flow-adjusted growth-sleeve unit value from its previous peak."),
    ]
    for term, definition in defs:
        p = set_para(doc.add_paragraph(), after=1, line=1.0)
        style_run(p.add_run(term + "  "), 7.5, True, NAVY)
        style_run(p.add_run(definition), 7.5, False, INK)

    doc.core_properties.title = "Kappa Foundation Initial Decision Record"
    doc.core_properties.subject = "Proposed initial investment allocation as of 31 December 2022"
    doc.core_properties.author = "Kappa Foundation independent investment team"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    make()
