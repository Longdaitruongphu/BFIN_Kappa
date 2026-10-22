"""Check the prepared 22 October package without fetching fresh data.

Run after run_initial.py. Optional Word page-count verification reads the
Microsoft Word export at outputs/initial_docx_word_render.pdf when present.
"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs/submission_readiness_20261021.json"


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    record = json.loads((ROOT / "outputs/initial_decision_rebuild.json").read_text())
    docx_path = ROOT / "deliverables/Kappa_Initial_Decision_Record_22Oct2026.docx"
    pdf_path = ROOT / "deliverables/Kappa_Initial_Decision_Record_22Oct2026.pdf"
    required = [
        pdf_path, docx_path, ROOT / "analysis/run_initial.py",
        ROOT / "analysis/retrieve_data.py", ROOT / "analysis/compare_initial_candidates.py",
        ROOT / "data/DATA_RECORD.md", ROOT / "data/raw/retrieval_manifest.jsonl",
        ROOT / "requirements-analysis.txt", ROOT / "requirements-document.txt",
        ROOT / "policy/initial_policy_proposed.json",
    ]
    missing = [str(p.relative_to(ROOT)) for p in required if not p.is_file()]
    if missing:
        raise FileNotFoundError(missing)
    raw = list((ROOT / "data/raw/train").glob("*"))
    assert len(list((ROOT / "data/raw/train").glob("*_daily.csv"))) == 10
    assert len(list((ROOT / "data/raw/train").glob("*.zip"))) == 1
    pdf_reader = PdfReader(pdf_path)
    assert len(pdf_reader.pages) == 2
    pdf_body = "\n".join(page.extract_text() or "" for page in pdf_reader.pages)
    for token in ("Full-universe maximum utility at A=4", "MSFT 2.92%", "USMV 46.18%",
                  "VTI 17.89%", "Max utility", "re-optimises all 10 candidates"):
        assert token in pdf_body, token
    docx = Document(docx_path)
    body = "\n".join([p.text for p in docx.paragraphs] +
                     [c.text for t in docx.tables for row in t.rows for c in row.cells])
    for token in ("6,346,350.71", "2,214,672.00", "3,429,808.70", "12,000,000.00",
                  "7.14%", "8.45%", "2.50%", "3.00%", "2023-Dec 2025",
                  "MSFT", "46.18%", "17.89%", "3.96%"):
        assert token in body, token
    optimizer = json.loads((ROOT / "outputs/sample_utility_robustness.json").read_text())
    weights = optimizer["weights_96m"]
    assert abs(sum(weights.values()) - 1.0) < 1e-10
    assert sum(w > 1e-10 for w in weights.values()) == 8
    assert abs(weights["USMV"] - 0.46175374093622124) < 1e-12
    assert abs(weights["VTI"] - 0.17888628710762142) < 1e-12
    pref = list(__import__("csv").DictReader((ROOT / "outputs/max_utility_preference.csv").open()))
    assert [int(row["risk_aversion_A"]) for row in pref] == [2, 4, 6, 8]
    assert abs(float(pref[1]["utility"]) - optimizer["utility_A4"]) < 1e-12
    budget = sum(record[k] for k in ("beta_initial_value", "delta_initial_value",
                                     "bond_fee", "growth_initial_value", "growth_fee"))
    assert abs(budget - 12e6) < 1e-6
    assert record["payment_2024_covered"] == 2.4e6
    assert record["payment_2027_covered"] == 3.6e6
    assert record["training_months"] == 96
    assert record["training_end"] == "2022-12"
    assert not (ROOT / "data/raw/evaluation").exists()
    manifest_lines = [json.loads(line) for line in
                      (ROOT / "data/raw/retrieval_manifest.jsonl").read_text().splitlines() if line]
    success = [row for row in manifest_lines if row.get("event") == "success"]
    assert len(success) == 11
    for row in success:
        source_file = ROOT / Path(row["file"].replace("\\", "/"))
        assert sha(source_file) == row["sha256"]
    with ZipFile(docx_path) as z:
        assert z.testzip() is None
        assert "word/document.xml" in z.namelist()
    word_render = ROOT / "outputs/initial_docx_word_render.pdf"
    if word_render.is_file():
        assert len(PdfReader(word_render).pages) == 2
    result = {
        "status": "PASS",
        "prepared_not_uploaded": True,
        "pdf_pages": 2,
        "word_export_pages": len(PdfReader(word_render).pages) if word_render.is_file() else None,
        "sources": "10 Yahoo CSV plus one Kenneth French ZIP, hashes checked",
        "training_months": 96,
        "max_utility_A4_nonzero_positions": 8,
        "max_utility_A4_weights_sum": round(sum(weights.values()), 12),
        "training_end": "2022-12",
        "budget_usd": round(budget, 2),
        "payment_2024_usd": record["payment_2024_covered"],
        "payment_2027_usd": record["payment_2027_covered"],
        "pdf_sha256": sha(pdf_path),
        "docx_sha256": sha(docx_path),
        "required_files_checked": len(required),
        "known_limits": [
            "SEC filing bytes were not saved after HTTP 403; dated eligibility references remain",
            "2026 vendor histories are retrospective and may differ from case-date displays",
            "No revised mandate was available for this original submission",
        ],
    }
    OUT.write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
